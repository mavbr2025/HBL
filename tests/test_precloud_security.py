import asyncio
import json
import os
from pathlib import Path
import subprocess
from urllib.parse import parse_qs, urlparse

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from mtm_hbl.api import main as api
from mtm_hbl.clickup_connector.oauth import ClickUpOAuthToken, OAuthStateStore
from mtm_hbl.clickup_hbl_generator import generate_hbl_from_clickup, _data_from_clickup_task
from mtm_hbl.config import Settings, get_settings
from mtm_hbl.excel.hbl_writer import ExcelHblWriter
from mtm_hbl.models.canonical import QaIssue
from mtm_hbl.models.clickup import ClickUpCustomField, ClickUpTaskData
from mtm_hbl.utils.file_naming import contained_output_path, safe_filename_component
from tests.conftest import valid_data
from tests.test_clickup_hbl_generator import FakeClickUpClient


@pytest.fixture
def api_client(monkeypatch, tmp_path):
    monkeypatch.setenv("HBL_API_TOKEN", "operator")
    monkeypatch.setenv("CLICKUP_CLIENT_ID", "client")
    monkeypatch.setenv("TOKEN_STORE_PATH", str(tmp_path / "token.json"))
    get_settings.cache_clear()
    monkeypatch.setattr(api, "state_store", OAuthStateStore())
    with TestClient(api.app) as client:
        yield client
    get_settings.cache_clear()


def test_every_administrative_route_requires_auth_before_validation(api_client, monkeypatch):
    public = {"/health", "/", "/auth/clickup/callback"}
    for route in api.app.routes:
        if route.path in public:
            continue
        for method in route.methods:
            path = route.path.replace("{task_id}", "task")
            assert api_client.request(method, path).status_code == 401
            assert api_client.request(method, path, headers={"Authorization": "Bearer wrong"}).status_code == 403
    assert api_client.get("/health").status_code == 200
    assert api_client.get("/auth/clickup/status", headers={"Authorization": "Bearer operator"}).status_code == 200
    monkeypatch.setenv("HBL_API_TOKEN", " ")
    get_settings.cache_clear()
    assert api_client.post("/packages/issue-dev", json={}).status_code == 503


@pytest.mark.parametrize("callback_path", ["/", "/auth/clickup/callback"])
def test_oauth_requires_initiated_browser_and_single_use_state(api_client, monkeypatch, callback_path):
    saved = []
    async def exchange(self, code):
        saved.append(code)
        return ClickUpOAuthToken(access_token="test")
    monkeypatch.setattr(api.ClickUpOAuthClient, "exchange_code", exchange)
    monkeypatch.setattr(api.LocalTokenStore, "save", lambda *args: None)
    assert api_client.get(callback_path, params={"code": "evil"}).status_code == 400
    started = api_client.get("/auth/clickup/start?redirect=false", headers={"Authorization": "Bearer operator"})
    state = parse_qs(urlparse(started.json()["authorization_url"]).query)["state"][0]
    assert "HttpOnly" in started.headers["set-cookie"]
    with TestClient(api.app) as other_browser:
        assert other_browser.get(callback_path, params={"code": "evil", "state": state}).status_code == 400
    assert api_client.get(callback_path, params={"code": "good", "state": state}).status_code == 200
    assert api_client.get(callback_path, params={"code": "replay", "state": state}).status_code == 400
    assert saved == ["good"]


def test_oauth_expired_state_is_rejected(monkeypatch):
    from mtm_hbl.clickup_connector import oauth
    clock = [0.0]
    monkeypatch.setattr(oauth.time, "monotonic", lambda: clock[0])
    store = OAuthStateStore()
    state = store.create("browser")
    clock[0] = 601
    assert store.consume(state, "browser") is False


@pytest.mark.parametrize("route", ["/packages/generate", "/packages/issue-dev"])
def test_raw_api_issuance_also_recalculates_qa_and_rejects_paths(api_client, route, tmp_path):
    data = valid_data()
    data.shipment.mbl_no = ""
    payload = {"review_packet": data.model_dump(mode="json"), "output_dir": str(tmp_path / "never-written")}
    headers = {"Authorization": "Bearer operator"}
    assert api_client.post(route, json=payload, headers=headers).status_code == 409
    data = valid_data()
    data.shipment.mtm_hbl_no = "X/../../outside"
    payload["review_packet"] = data.model_dump(mode="json")
    assert api_client.post(route, json=payload, headers=headers).status_code == 400
    assert not (tmp_path / "never-written").exists()


def approved_task(data):
    return ClickUpTaskData(id="task", custom_fields=[
        ClickUpCustomField(id="json", name="Canonical HBL JSON", value=data.model_dump_json()),
        ClickUpCustomField(id="approval", name="HBL Approval Status", value="Approved"),
        ClickUpCustomField(id="by", name="HBL Approved By", value="Operator"),
        ClickUpCustomField(id="at", name="HBL Approved At", value="2026-09-05"),
    ])


@pytest.mark.parametrize("corruption", ["missing-mbl", "missing-packages", "bad-country", "missing-live-hbl"])
def test_canonical_self_reported_qa_cannot_authorize_issuance(tmp_path, app_config, corruption):
    data = valid_data()
    live = {"hbl_number": data.shipment.mtm_hbl_no, "owner_country": "Guatemala"}
    if corruption == "missing-mbl":
        data.shipment.mbl_no = ""
    elif corruption == "missing-packages":
        data.containers[0].package_count = ""
        data.qa.soft_warnings.append(QaIssue(id="package_counts_total_only", severity="soft_warning", field="containers", message="fake", blocking_scope="none"))
    elif corruption == "bad-country":
        live["owner_country"] = ""
    else:
        live["hbl_number"] = ""
    data.qa.hard_errors.clear()
    client = FakeClickUpClient(approved_task(data), live)
    with pytest.raises(ValueError, match="hard QA"):
        asyncio.run(generate_hbl_from_clickup(task_ref="task", client=client, settings=Settings(runs_dir=tmp_path),
                    app_config=app_config, mode="issue"))
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("identifier", ["X/../../escape", "X\\..\\escape", ".", "..", "C:escape", "bad\nname"])
def test_unsafe_identifiers_are_rejected(identifier):
    with pytest.raises(ValueError):
        safe_filename_component(identifier)


def test_output_symlink_cannot_escape(tmp_path):
    base = tmp_path / "base"
    base.mkdir()
    (base / "report.json").symlink_to(tmp_path / "outside.json")
    with pytest.raises(ValueError):
        contained_output_path(base, "report.json")


def test_workbook_external_text_is_literal_and_template_formula_survives(tmp_path, app_config):
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    workbook.active["L55"] = "=D55*C72"
    workbook.save(template)
    data = valid_data()
    data.parties.shipper.raw_text = "=1+1\n=2+2"
    data.cargo.description_raw = "=3+3"
    data.containers[0].seal_no = "=4+4"
    output = tmp_path / "out.xlsx"
    ExcelHblWriter(app_config).write(template, output, data)
    sheet = load_workbook(output).active
    for value in ["=1+1", "=2+2", "=3+3", "=4+4"]:
        matches = [cell for row in sheet for cell in row if cell.value == value]
        assert matches and all(cell.data_type == "s" for cell in matches)
    assert sheet["L55"].data_type == "f"


def test_normal_deployer_cannot_modify_role_permissions():
    path = Path(__file__).parents[1] / "aws/verification-service/deploy-iam-policy.json"
    policy = json.loads(path.read_text())
    iam = [entry for entry in policy["Statement"] if any(action.startswith("iam:") for action in entry["Action"])]
    assert {action for entry in iam for action in entry["Action"]} == {"iam:GetRole", "iam:PassRole"}
    pass_role = next(entry for entry in iam if "iam:PassRole" in entry["Action"])
    assert pass_role["Condition"]["StringEquals"]["iam:PassedToService"] == "lambda.amazonaws.com"


def test_missing_preprovisioned_role_fails_before_aws_writes(tmp_path):
    fake_aws = tmp_path / "aws"
    log = tmp_path / "calls"
    fake_aws.write_text('#!/bin/bash\nprintf "%s\\n" "$*" >> "$CALL_LOG"\ncase "$1 $2" in\n "configure get") echo us-east-1;;\n "sts get-caller-identity") echo 123456789012;;\n "iam get-role") exit 1;;\n *) exit 99;;\nesac\n')
    fake_aws.chmod(0o755)
    script = Path(__file__).parents[1] / "aws/verification-service/scripts/deploy_aws_cli.sh"
    result = subprocess.run(["bash", str(script)], env={**os.environ, "PATH": str(tmp_path) + os.pathsep + os.environ["PATH"], "CALL_LOG": str(log)}, capture_output=True)
    assert result.returncode != 0
    assert all(line.startswith(("configure get", "sts get-caller-identity", "iam get-role")) for line in log.read_text().splitlines())


def test_approved_total_only_exception_is_rederived_with_exact_cargo_binding(app_config):
    from mtm_hbl.resolver.customer_rules import apply_customer_profile, restore_trusted_package_exception
    from mtm_hbl.validation.validation_engine import ValidationEngine
    data = valid_data()
    data.shipment.clickup_task_id = "MTMLXGT-25972"
    data.shipment.mtm_hbl_no = "GOSZX26041381"
    apply_customer_profile(data, app_config, "repuestos_acquaroni")
    data.qa = type(data.qa)()
    restore_trusted_package_exception(data, app_config)
    ValidationEngine(app_config).validate(data)
    assert any(issue.id == "package_counts_total_only" for issue in data.qa.soft_warnings)
    assert not data.qa.hard_errors
    for field in ("hbl", "total", "container"):
        altered = data.model_copy(deep=True)
        altered.qa = type(altered.qa)()
        if field == "hbl":
            altered.shipment.mtm_hbl_no = "OTHER"
        elif field == "total":
            altered.cargo.total_packages = "999"
        else:
            altered.containers[0].container_no = "OTHER"
        restore_trusted_package_exception(altered, app_config)
        ValidationEngine(app_config).validate(altered)
        assert not any(issue.id == "package_counts_total_only" for issue in altered.qa.soft_warnings)
        assert altered.qa.hard_errors
